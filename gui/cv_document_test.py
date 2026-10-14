import unittest, tempfile, hashlib, zipfile
from pathlib import Path
from docx import Document
from lxml import etree as ET
from cv_document import main, NS

class DocumentTests(unittest.TestCase):
    def test_highlights_cross_run_keywords_and_strips_without_accepting_revisions(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);master=root/'master.docx';doc=Document();p=doc.add_paragraph();p.add_run('Py');p.add_run('thon').bold=True;p.add_run(' and machine learning. Python again.');doc.save(master)
            digest=hashlib.sha256(master.read_bytes()).hexdigest();before='Python and machine learning. Python again.'
            cross_run=main({'operation':'export','master':str(master),'master_sha256':digest,'changes':[],'keywords':['Python','machine learning'],'output':str(root/'cross-run')})
            self.assertEqual(cross_run['highlighted']['keyword_occurrences'],3);self.assertEqual(cross_run['highlighted']['highlight_runs'],4)
            result=main({'operation':'export','master':str(master),'master_sha256':digest,'changes':[{'index':0,'before':before,'after':'Python machine learning engineer.','accepted':True}],'keywords':['Python','machine learning','Kubernetes'],'output':str(root/'exports')})
            self.assertGreater(result['highlighted']['highlight_runs'],0);self.assertEqual(result['tracked']['highlight_runs'],0);self.assertEqual(result['final']['highlight_runs'],0)
            highlighted=root/'exports/cv_tracked_highlighted.docx';stripped=root/'stripped.docx'
            stripped_info=main({'operation':'strip_highlights','master':str(highlighted),'output':str(stripped)})
            self.assertGreater(stripped_info['removed_highlights'],0);self.assertEqual(stripped_info['remaining_revisions'],2)
            with zipfile.ZipFile(stripped) as z:
                tree=ET.fromstring(z.read('word/document.xml'));self.assertEqual(len(tree.findall('.//w:highlight',NS)),0);self.assertEqual(len(tree.findall('.//w:ins',NS)),1)
            accepted_info=main({'operation':'strip_highlights','master':str(root/'exports/cv_final.docx'),'output':str(root/'final-stripped.docx')})
            self.assertEqual(accepted_info['remaining_revisions'],0);self.assertEqual(Document(root/'final-stripped.docx').paragraphs[0].text,'Python machine learning engineer.')
    def test_export_accept_reject_edit_and_preserve_master(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);master=root/'master.docx';doc=Document();doc.add_paragraph('Original supported claim');doc.add_paragraph('Leave this unchanged');doc.save(master)
            digest=hashlib.sha256(master.read_bytes()).hexdigest()
            changes=[{'index':0,'before':'Original supported claim','after':'Edited supported claim','accepted':True},{'index':1,'before':'Leave this unchanged','after':'Rejected claim','accepted':False}]
            result=main({'operation':'export','master':str(master),'master_sha256':digest,'changes':changes,'output':str(root/'exports')})
            final=Document(root/'exports/cv_final.docx');self.assertEqual([p.text for p in final.paragraphs],['Edited supported claim','Leave this unchanged'])
            self.assertEqual(result['tracked']['insertions'],1);self.assertEqual(result['final']['insertions'],0)
            self.assertEqual(hashlib.sha256(master.read_bytes()).hexdigest(),digest)
            with zipfile.ZipFile(root/'exports/cv_tracked.docx') as z:
                tree=ET.fromstring(z.read('word/document.xml'));self.assertEqual(tree.find('.//w:delText',NS).text,'Original supported claim');self.assertEqual(tree.find('.//w:ins/w:r/w:t',NS).text,'Edited supported claim')
            with zipfile.ZipFile(master) as original,zipfile.ZipFile(root/'exports/cv_final.docx') as output:
                for name in original.namelist():
                    if name not in ['word/document.xml','word/settings.xml']:self.assertEqual(original.read(name),output.read(name))
            with self.assertRaisesRegex(ValueError,'Master CV changed'):
                main({'operation':'export','master':str(master),'master_sha256':'wrong','changes':changes,'output':str(root/'bad')})

if __name__=='__main__':unittest.main()
